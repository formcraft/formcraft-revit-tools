# Bar marks on the drawings: every member's BBS bar marks (Mark) and its reinforcement (Comments) written into the
# Revit model, and itemised reinforcement schedules (columns, beams, foundations) on sheet S-110, so the BBS and the
# drawings read together. fmc_model_edit ops; then revit-edit.mjs them.
#   python scripts/structure/bar-marks.py <work dir> <BBS xlsx> <rc-design json>   (openpyxl: not -I)
import json, sys, re
import openpyxl
W, X, DJ = sys.argv[1:4]
d = json.load(open(DJ, encoding="utf-8"))
marks = {}
for ws in openpyxl.load_workbook(X, read_only=True).worksheets[1:]:
    for r in list(ws.iter_rows(values_only=True))[1:]:
        m = re.match(r"(?:GB|C|B|F)(\d+)\b", str(r[0] or ""))
        if m: marks.setdefault(int(m.group(1)), []).append(r[1])
rein = {c["id"]: f'{c["bars"]}, links {c["links"]}' for c in d["columns"]}
rein |= {b["id"]: f'B {b["bottom"]}, T {b["top"]}, links {b["links"]}' for b in d["beams"] + d["ground_beams"]}
rein |= {p["id"]: f'{p["size"]}, {p["bars"]}' for p in d["pads"]}
ops = []; n = 0
def add(o):
    global n; ops.append(o); n += 1; return f"${n}"
for i, ms in sorted(marks.items()):
    if i not in rein: continue
    add({"op": "set_param", "ids": [i], "param": "Mark", "value_text": f"{ms[0]}–{ms[-1]}" if len(ms) > 1 else ms[0]})
    add({"op": "set_param", "ids": [i], "param": "Comments", "value_text": rein[i]})
sch = []
for cat, name, fields, sort in (("Structural Columns", "S Column reinforcement", ["Mark", "Type", "Base Level", "Comments"], ["Base Level", "Mark"]),
                                ("Structural Framing", "S Beam reinforcement", ["Mark", "Type", "Reference Level", "Comments"], ["Reference Level", "Mark"]),
                                ("Structural Foundations", "S Foundation reinforcement", ["Mark", "Type", "Comments"], ["Mark"])):
    sch.append(add({"op": "create_schedule", "category": cat, "name": name, "fields": fields, "sort": sort, "itemize": True}))
sh = add({"op": "create_sheet", "number": "S-110", "name": "Reinforcement schedules (bar marks as BBS) - PRELIMINARY, NOT FOR CONSTRUCTION"})
for k, s in enumerate(sch): add({"op": "place_view", "sheet": sh, "view": s, "point_mm": [30 + k * 270, 560]})
json.dump(ops, open(W + "/marks_ops.json", "w"), indent=0)
print(len(ops), "ops;", sum(1 for i in marks if i in rein), "members marked")
